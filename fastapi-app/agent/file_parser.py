"""上传附件解析：txt/md/pdf/xlsx 文本提取 + 图片识图（视觉模型或人工）。"""
import os

from models import Analysis, FileAsset


def _extract_xlsx(path: str) -> str:
    """把 Excel 表格转成 "表头=值" 的文本行。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0] if rows else []
        for row in rows[1:]:
            pairs = [f"{header[i] or '列'+str(i)}={row[i]}" for i in range(len(row)) if row[i] is not None]
            if pairs:
                lines.append("；".join(pairs))
    wb.close()
    return "\n".join(lines)


def _extract_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def extract_file_text(asset: FileAsset) -> str:
    """按文件类型提取文本；无法提取返回空串。"""
    path = asset.file_path
    if not os.path.isfile(path):
        return ""
    try:
        if asset.file_type == "excel":
            return _extract_xlsx(path)
        if asset.file_type == "pdf":
            return _extract_pdf(path)
        if asset.file_type in ("text",):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    except Exception:
        return ""
    return ""


async def extract_analysis_files(analysis_id: int) -> dict:
    """读取分析的上传附件，返回 {'text': str, 'unrecognized_images': int}。

    text：txt/md/pdf/xlsx 提取的文本 + 识图成功后的描述。
    unrecognized_images：无视觉模型时未能识别的图片数（需人工补充）。
    """
    analysis = await Analysis.filter(id=analysis_id).first()
    if not analysis or not analysis.input_files:
        return {"text": "", "unrecognized_images": 0}

    from agent import multimodal
    text_parts, unrecognized = [], 0
    assets = await FileAsset.filter(id__in=analysis.input_files)
    for asset in assets:
        if asset.file_type == "image":
            desc = await multimodal.describe_image(asset.file_path)
            if desc:
                text_parts.append(f"结构草图描述：{desc}")
            else:
                unrecognized += 1
        else:
            txt = extract_file_text(asset)
            if txt:
                text_parts.append(txt)
    return {"text": "\n".join(text_parts), "unrecognized_images": unrecognized}
